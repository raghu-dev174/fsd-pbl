import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = REPORT_DIR / "FSD_Internal_2_Real_Time_Transport_Report.docx"
SUMMARY_PATH = ROOT / "results" / "summary.json"
AUTOCANNON_PATH = ROOT / "results" / "autocannon-short-polling.json"

with SUMMARY_PATH.open("r", encoding="utf-8") as handle:
    summary = json.load(handle)
with AUTOCANNON_PATH.open("r", encoding="utf-8") as handle:
    autocannon = json.load(handle)

short = summary["transports"]["shortPolling"]
long_poll = summary["transports"]["longPolling"]
sse = summary["transports"]["sse"]

NAVY = "172554"
BLUE = "2563EB"
PURPLE = "7C3AED"
GREEN = "059669"
PALE_BLUE = "EFF6FF"
PALE_GRAY = "F8FAFC"
BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)

def font_file(size, bold=False):
    names = ["arialbd.ttf" if bold else "arial.ttf", "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]
    for name in names:
        for base in [Path("C:/Windows/Fonts"), Path("/usr/share/fonts/truetype/dejavu")]:
            candidate = base / name
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color=BORDER, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))

def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_repeat_table_layout(table, widths=None, numeric_columns=()):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for r_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if r_index == 0:
            repeat_table_header(row)
        for c_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[c_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    if r_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if r_index == 0:
                set_cell_shading(cell, NAVY)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)

def add_table(doc, headers, rows, widths=None, numeric_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
    set_repeat_table_layout(table, widths, numeric_columns)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])

def add_caption(doc, text):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = False
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = MUTED
    return paragraph

def add_image(doc, image_path, caption, max_width=6.7, max_height=7.2, page_break=False):
    if page_break:
        doc.add_page_break()
    with Image.open(image_path) as img:
        width_px, height_px = img.size
    aspect = width_px / height_px
    width = min(max_width, max_height * aspect)
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    add_caption(doc, caption)

def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)

def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)

def add_manual_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        number_run = paragraph.add_run(f"{index}. ")
        number_run.bold = True
        paragraph.add_run(item)

def add_section(doc, number, title, page_break=False):
    if page_break:
        doc.add_page_break()
    heading = doc.add_heading(f"{number} {title}", level=1)
    heading.paragraph_format.keep_with_next = True
    return heading

def add_subheading(doc, title):
    heading = doc.add_heading(title, level=2)
    heading.paragraph_format.keep_with_next = True
    return heading

def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        run.bold = True
    paragraph.add_run(text)
    return paragraph

def create_architecture():
    path = ROOT / "charts" / "system-architecture.png"
    width, height = 1500, 780
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = font_file(42, True)
    box_font = font_file(27, True)
    small_font = font_file(22)
    draw.text((70, 55), "System Architecture", font=title_font, fill="#0F172A")
    draw.text((70, 112), "One source feeds all transports so the comparison uses the same events.", font=small_font, fill="#64748B")

    def box(coords, heading, detail, fill, outline):
        draw.rounded_rectangle(coords, radius=22, fill=fill, outline=outline, width=4)
        x0, y0, x1, y1 = coords
        hb = draw.textbbox((0, 0), heading, font=box_font)
        db = draw.textbbox((0, 0), detail, font=small_font)
        draw.text(((x0 + x1 - (hb[2] - hb[0])) / 2, y0 + 32), heading, font=box_font, fill="#0F172A")
        draw.text(((x0 + x1 - (db[2] - db[0])) / 2, y0 + 85), detail, font=small_font, fill="#475569")

    box((480, 180, 1020, 330), "Shared Notification Source", "ID + message + generatedAt every 2 seconds", "#EFF6FF", "#2563EB")
    boxes = [
        ((60, 460, 440, 640), "Short Polling", "GET every 5 seconds", "#EFF6FF", "#2563EB"),
        ((560, 460, 940, 640), "Long Polling", "Held until event or timeout", "#F5F3FF", "#7C3AED"),
        ((1060, 460, 1440, 640), "Server-Sent Events", "Persistent text/event-stream", "#ECFDF5", "#059669"),
    ]
    for args in boxes:
        box(*args)
    source_bottom = (750, 330)
    for x in (250, 750, 1250):
        draw.line((source_bottom[0], source_bottom[1], x, 440), fill="#64748B", width=5)
        draw.polygon([(x, 452), (x - 11, 432), (x + 11, 432)], fill="#64748B")
    draw.text((440, 700), "React UI + 100-client concurrent simulator", font=box_font, fill="#0F172A")
    image.save(path, optimize=True)
    return path

architecture_path = create_architecture()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

for name, size, spacing in (("Title", 28, 12), ("Subtitle", 15, 10), ("Heading 1", 18, 10), ("Heading 2", 13, 7)):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(spacing)
    style.paragraph_format.space_after = Pt(spacing / 2)
    style.paragraph_format.keep_with_next = True

title_p_pr = styles["Title"]._element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

styles["Caption"].font.name = "Arial"
styles["Caption"].font.color.rgb = MUTED

for sec in doc.sections:
    footer = sec.footer
    footer_p = footer.paragraphs[0]
    footer_p.add_run("FSD Internal 2  |  Assignment 19                                      ")
    add_page_number(footer_p)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("FSD Internal 2")
direct_title_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if direct_title_border is not None:
    title._p.get_or_add_pPr().remove(direct_title_border)
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("Polling vs Long Polling vs Server Sent Events")
lead = doc.add_paragraph()
lead.paragraph_format.space_before = Pt(18)
lead_run = lead.add_run("Implementation and measured comparison of real-time notification transports")
lead_run.font.size = Pt(16)
lead_run.font.bold = True
lead_run.font.color.rgb = TEXT

doc.add_paragraph("Assignment Number 19")
doc.add_paragraph("Subject React")
doc.add_paragraph("Topic Real-time transport trade-offs")
doc.add_paragraph("Benchmark 100 concurrent clients, 120 seconds per transport")
doc.add_paragraph("Date 7 October 2026")

doc.add_paragraph().paragraph_format.space_after = Pt(16)
cover_summary = doc.add_paragraph()
cover_summary.add_run("Main finding. ").bold = True
cover_summary.add_run(
    "SSE provided the lowest measured request rate, average CPU use, and delivery delay for this one-way notification workload. "
    "Short polling remained the simplest behaviorally but introduced multi-second waiting. Long polling reduced delay, yet its request rate rose with the event frequency because every event completed and restarted 100 requests."
)

doc.add_page_break()
doc.add_heading("Report Contents", level=1)
contents = [
    "1 Assignment Details", "2 Objective", "3 Tools Used", "4 Introduction", "5 System Architecture",
    "6 Implementation", "7 100 Client Simulation", "8 Benchmark Methodology", "9 Screenshots",
    "10 Results", "11 Charts", "12 Analysis", "13 Advantages and Disadvantages",
    "14 When Should Each Be Used", "15 Observations", "16 Conclusion", "17 Future Improvements"
]
add_manual_numbered(doc, [item.split(' ', 1)[1] for item in contents])

add_section(doc, 1, "Assignment Details")
add_table(doc, ["Field", "Value"], [
    ["Assignment Number", "19"], ["Subject", "React"],
    ["Title", "Polling vs Long Polling vs Server-Sent Events"],
    ["Topic", "Real-time transport trade-offs"],
    ["Required comparison", "100 concurrent clients using a common notification workload"]
], widths=[2.0, 4.8])

add_section(doc, 2, "Objective")
add_body(doc, "The objective was to implement one live notification feed through three HTTP delivery strategies, measure their behavior under the same 100-client workload, and use the evidence to choose the least complex transport that still satisfies delivery requirements. The project measures request volume, Node server CPU, connection behavior, and end-to-end delivery delay.")

add_section(doc, 3, "Tools Used")
add_table(doc, ["Tool", "Use in this project"], [
    ["React", "Interactive demonstration UI with separate transport views and measured results"],
    ["Node.js 24.18.0", "HTTP server, shared notification source, metrics, and concurrent client simulator"],
    ["Express", "API routes, static React build, and SSE response handling"],
    ["EventSource API", "Browser SSE connection in the React client"],
    ["autocannon", "Separate HTTP request-response capacity check for the short-poll endpoint"],
    ["JavaScript", "Server, frontend, simulation, statistics, and validation"],
    ["process.cpuUsage", "One-second Node process CPU sampling"],
    ["Browser capture", "Screenshots of the running UI and saved benchmark evidence"],
    ["Python Pillow and python-docx", "Measured-data charts and final report generation"]
], widths=[1.7, 5.1])

add_section(doc, 4, "Introduction")
add_body(doc, "Real-time web communication describes how a browser learns about server-side changes without requiring a manual refresh. The transport choice affects request overhead, delivery delay, connection lifetime, proxy behavior, browser support, and implementation complexity.")
add_subheading(doc, "Short Polling")
add_body(doc, "The client sends a normal HTTP request on a fixed schedule. This is easy to implement and debug, but most requests may return no new data. Delivery can be delayed by almost the full polling interval.")
add_subheading(doc, "Long Polling")
add_body(doc, "The server holds a request until data is available or a timeout occurs. The client reconnects immediately after each response. This provides near-real-time delivery over ordinary HTTP, but every event completes a response and creates another request for each client.")
add_subheading(doc, "Server-Sent Events")
add_body(doc, "SSE keeps one HTTP response open with the text/event-stream content type. The server pushes one-way events to the browser, and the browser EventSource API handles the stream and reconnection behavior. SSE is a natural fit for notifications, dashboards, and live feeds.")
add_body(doc, "Transport choice matters because the simplest solution is not always the one with the fewest lines of code or the smallest delay. A fixed interval may be acceptable for infrequent, low-priority changes. A persistent stream is more efficient when updates must arrive quickly and the data flow is primarily server-to-client.")

add_section(doc, 5, "System Architecture", page_break=True)
add_image(doc, architecture_path, "Figure 1. Common notification source and the three delivery paths", max_width=6.8, max_height=4.1)
add_body(doc, "One NotificationSource object generates an event every two seconds. Each event has a monotonically increasing ID, message, generatedAt millisecond timestamp, and ISO timestamp. The same object supplies the short-poll response, releases all eligible long-poll waiters, and writes the event to every SSE stream. This prevents one transport from receiving an easier or different workload.")

add_section(doc, 6, "Implementation")
add_subheading(doc, "Short Polling")
add_body(doc, "The GET /api/poll endpoint returns every stored notification whose ID is greater than the client's since value. The React client starts one request immediately, records each notification, waits until five seconds have elapsed since the cycle began, and then sends the next request. The server increments its poll request counter on entry and records how many notifications were returned.")
add_subheading(doc, "Long Polling")
add_body(doc, "The GET /api/long-poll endpoint first checks for a backlog. If none is available, the response is stored in an in-memory waiter set and remains open. A generated notification completes every eligible waiter. A 15-second timeout prevents an indefinitely orphaned request, and the close handler removes disconnected clients. The browser and simulator reconnect immediately after a response.")
add_subheading(doc, "Server-Sent Events")
add_body(doc, "The GET /api/events endpoint sets Content-Type to text/event-stream, disables caching and buffering, flushes the headers, and sends named notification events with SSE IDs. A comment heartbeat is written every 10 seconds. Disconnect handlers clear the heartbeat, remove the client from the connection map, and record stream duration. The React UI uses new EventSource('/api/events').")
add_subheading(doc, "Instrumentation")
add_body(doc, "The server counts requests by transport, generated and dispatched notifications, active and peak connections, and connection durations. CPU samples are taken once per second. The simulator stores one raw delivery record per client and notification, including client ID, notification ID, generation time, receipt time, and calculated delay.")

add_section(doc, 7, "100 Client Simulation")
add_body(doc, "The benchmark creates an array of 100 asynchronous client tasks and starts them together with Promise.all. Each task has a unique, zero-padded client ID. Short-poll clients execute independent five-second cycles. Long-poll clients maintain one outstanding request and reconnect as soon as it resolves. SSE clients maintain 100 simultaneous streaming fetch responses and parse SSE frames. The production React view uses the browser EventSource API; the Node simulator uses streaming fetch so receipt timestamps can be collected uniformly in one process.")
add_body(doc, "All transports were tested with a fresh server process, a reset notification sequence, the same two-second generator interval, and the same 120-second target duration. A 10-second no-client phase preceded each run to capture an idle baseline with the generator still active.")

add_section(doc, 8, "Benchmark Methodology")
add_table(doc, ["Setting", "Controlled value"], [
    ["Concurrent clients", "100"], ["Measured duration", "120 seconds per transport"],
    ["Idle baseline", "10 seconds before each transport"], ["Notification interval", "2,000 ms"],
    ["Short-poll interval", "5,000 ms"], ["Long-poll timeout", "15,000 ms"],
    ["Environment", "Windows NT 10.0.26200.0, local loopback, Node.js 24.18.0"],
    ["Run order", "Short polling, long polling, SSE"]
], widths=[2.2, 4.6])
add_subheading(doc, "Request Counting")
add_body(doc, "Every transport endpoint increments a server-side counter when a request arrives. Requests per minute equals the recorded count divided by the actual measured duration in minutes. Values were not calculated from configuration assumptions.")
add_subheading(doc, "CPU Measurement")
add_body(doc, "The server samples process.cpuUsage once per second. For each sample, user plus system CPU microseconds are divided by elapsed wall-clock microseconds and multiplied by 100. The result is percent of one logical CPU core. Average and peak values use the same method for all transports. The baseline uses the same generator and sampler without simulated clients.")
add_subheading(doc, "Delivery Delay")
add_body(doc, "Each notification carries generatedAt from the server. On receipt, the simulator records Date.now as receivedAt. Delivery delay is receivedAt minus generatedAt in milliseconds. Average, median, minimum, maximum, and p95 values are calculated from every retained delivery sample.")
add_subheading(doc, "Role of autocannon")
add_body(doc, f"Autocannon was used for a separate 100-connection, 10-second capacity check of the ordinary short-poll endpoint. It measured {autocannon['requests']['average']:,.1f} requests/second, {autocannon['latency']['average']:.1f} ms average latency, and {autocannon['latency']['p99']:.0f} ms p99 latency. These values are not used in the three-transport comparison because autocannon does not reproduce held long-poll lifecycles or client-side SSE event timing.")

add_section(doc, 9, "Screenshots", page_break=True)
screens = [
    ("01_project_structure.jpg", "Figure 2. Implemented project structure and saved evidence"),
    ("02_react_live_feed.jpg", "Figure 3. React live notification demonstration"),
    ("03_short_polling.jpg", "Figure 4. Short-polling view with five-second request cycles"),
    ("04_long_polling.jpg", "Figure 5. Long-polling view while a request waits for the next event"),
    ("05_sse.jpg", "Figure 6. SSE view using one browser EventSource request"),
    ("06_100_clients.jpg", "Figure 7. Measured 100-client comparison dashboard"),
    ("07_autocannon_results.jpg", "Figure 8. Actual autocannon capacity-check output"),
    ("08_cpu_measurement.jpg", "Figure 9. Browser-rendered CPU measurement chart"),
    ("09_delivery_delay_results.jpg", "Figure 10. Browser-rendered delivery-delay chart"),
    ("10_comparison_table.jpg", "Figure 11. Actual comparison table from the result dashboard"),
    ("11_charts.jpg", "Figure 12. Browser-rendered measured charts")
]
for index, (filename, caption) in enumerate(screens):
    add_image(doc, ROOT / "screenshots" / filename, caption, max_width=6.55, max_height=7.65, page_break=index > 0)

add_section(doc, 10, "Results", page_break=True)
rows = [
    ["Total requests", f"{short['totalRequests']:,}", f"{long_poll['totalRequests']:,}", f"{sse['totalRequests']:,}"],
    ["Requests per minute", f"{short['requestsPerMinute']:,.1f}", f"{long_poll['requestsPerMinute']:,.1f}", f"{sse['requestsPerMinute']:,.1f}"],
    ["Average CPU", f"{short['averageCpuPercent']:.2f}%", f"{long_poll['averageCpuPercent']:.2f}%", f"{sse['averageCpuPercent']:.2f}%"],
    ["Peak CPU", f"{short['peakCpuPercent']:.2f}%", f"{long_poll['peakCpuPercent']:.2f}%", f"{sse['peakCpuPercent']:.2f}%"],
    ["Idle baseline CPU", f"{short['baselineAverageCpuPercent']:.2f}%", f"{long_poll['baselineAverageCpuPercent']:.2f}%", f"{sse['baselineAverageCpuPercent']:.2f}%"],
    ["Average delivery delay", f"{short['deliveryDelayMs']['average']:,.1f} ms", f"{long_poll['deliveryDelayMs']['average']:,.1f} ms", f"{sse['deliveryDelayMs']['average']:,.1f} ms"],
    ["Median delivery delay", f"{short['deliveryDelayMs']['median']:,.1f} ms", f"{long_poll['deliveryDelayMs']['median']:,.1f} ms", f"{sse['deliveryDelayMs']['median']:,.1f} ms"],
    ["Minimum delivery delay", f"{short['deliveryDelayMs']['min']} ms", f"{long_poll['deliveryDelayMs']['min']} ms", f"{sse['deliveryDelayMs']['min']} ms"],
    ["Maximum delivery delay", f"{short['deliveryDelayMs']['max']:,} ms", f"{long_poll['deliveryDelayMs']['max']:,} ms", f"{sse['deliveryDelayMs']['max']:,} ms"],
    ["p95 delivery delay", f"{short['deliveryDelayMs']['p95']:,} ms", f"{long_poll['deliveryDelayMs']['p95']:,} ms", f"{sse['deliveryDelayMs']['p95']:,} ms"],
    ["Delivery samples", f"{short['deliverySamples']:,}", f"{long_poll['deliverySamples']:,}", f"{sse['deliverySamples']:,}"],
    ["Peak active connections", f"{short['peakActiveConnections']}", f"{long_poll['peakActiveConnections']}", f"{sse['peakActiveConnections']}"]
]
add_table(doc, ["Metric", "Short Polling", "Long Polling", "SSE"], rows, widths=[2.35, 1.48, 1.48, 1.28], numeric_columns=(1, 2, 3))
add_body(doc, "The short-poll request rate matches the expected order of 100 clients x 12 requests per minute. Long polling produced more requests in this experiment because the notification interval was two seconds: each event completed 100 requests and caused 100 immediate reconnects. SSE opened 100 connections once, so its normalized rate was about 50 requests per minute over a two-minute test.")

add_section(doc, 11, "Charts", page_break=True)
add_image(doc, ROOT / "charts" / "requests-per-minute.png", "Figure 13. Measured requests per minute", max_width=6.6, max_height=4.2)
add_image(doc, ROOT / "charts" / "cpu-usage.png", "Figure 14. Average Node server CPU", max_width=6.6, max_height=4.2, page_break=True)
add_image(doc, ROOT / "charts" / "delivery-delay.png", "Figure 15. Average delivery delay", max_width=6.6, max_height=4.2, page_break=True)
add_image(doc, ROOT / "charts" / "overall-comparison.png", "Figure 16. Overall measured comparison", max_width=6.7, max_height=4.6, page_break=True)

add_section(doc, 12, "Analysis", page_break=True)
add_subheading(doc, "Request Overhead")
add_body(doc, "Short polling created a predictable 1,199.8 requests per minute. Some responses contained multiple notifications, and requests sent before an update returned no new data. Long polling created 2,992.2 requests per minute in this specific workload. That result is not a contradiction: long-poll traffic follows the event rate, and a two-second event interval is faster than the five-second polling interval. SSE required only the original stream requests unless a reconnect was needed.")
add_subheading(doc, "Persistent Connections")
add_body(doc, "Long polling and SSE both reached 100 active connections. The difference is lifecycle. A long-poll connection lasted about 1,955 ms on average and was replaced after each event. An SSE connection remained open for about 120 seconds. Persistent connections require server and proxy capacity, but SSE avoids repeated HTTP parsing and connection turnover.")
add_subheading(doc, "CPU Usage")
add_body(doc, "Average process CPU was 2.05% for short polling, 2.57% for long polling, and 0.87% for SSE. The SSE average was below its own 1.21% idle baseline, showing the small scale and ordinary timing noise of a single local run. The result still indicates that maintaining 100 idle streams was inexpensive for this Node process. Peak samples were bursty and should not be interpreted without the average and raw time series.")
add_subheading(doc, "Delivery Latency")
add_body(doc, "Short polling averaged 2,401.8 ms and reached 5,032 ms, which is consistent with events waiting for the next five-second request. Long polling reduced the average to 39.0 ms because events released waiting responses. SSE achieved 6.4 ms average delay because the server wrote directly to an already-open stream.")
add_subheading(doc, "Scalability and Complexity")
add_body(doc, "Short polling is simple but can waste requests as client counts grow. Long polling provides broad HTTP compatibility but requires waiter storage, timeouts, disconnect cleanup, and repeated reconnects. SSE requires correct streaming headers and proxy configuration, yet its browser API and one-way semantics are simpler than a bidirectional socket protocol for notification-only applications.")

add_section(doc, 13, "Advantages and Disadvantages")
add_table(doc, ["Approach", "Advantages", "Disadvantages"], [
    ["Short Polling", "Simple request-response behavior; easy caching, logging, and debugging; broadly compatible", "Unnecessary requests when no update exists; delivery waits for the next interval; traffic scales with clients and frequency"],
    ["Long Polling", "Near-real-time delivery over ordinary HTTP; works where streaming may be unsuitable", "Held-request management, timeouts, and disconnect cleanup; every event completes and restarts one request per client"],
    ["SSE", "Efficient one-way streaming; browser EventSource support; low measured delay and request churn", "Persistent connections; one-way only; proxies must avoid buffering and use appropriate timeouts"]
], widths=[1.25, 2.72, 2.72])

add_section(doc, 14, "When Should Each Be Used")
add_subheading(doc, "Short Polling")
add_bullets(doc, ["Infrequent updates where several seconds of delay is acceptable", "Extremely simple systems or operational environments", "Small client counts and APIs that already expose snapshot data"])
add_subheading(doc, "Long Polling")
add_bullets(doc, ["Older systems that need near-real-time behavior over familiar HTTP request-response infrastructure", "Networks or intermediaries where a continuous event stream is unsuitable", "Event rates low enough that reconnect churn remains reasonable"])
add_subheading(doc, "Server-Sent Events")
add_bullets(doc, ["Notifications, monitoring dashboards, activity streams, progress updates, and live feeds", "Applications whose real-time traffic is primarily server-to-browser", "Cases that benefit from automatic browser reconnection and event IDs"])
add_body(doc, "WebSockets may be preferable when the browser and server must exchange frequent real-time messages in both directions. WebSockets are outside this assignment and should not be selected merely because they are more general.")

add_section(doc, 15, "Observations")
add_manual_numbered(doc, [
    "The observed short-poll rate of 1,199.8 requests per minute was effectively the expected 1,200 requests per minute for 100 clients at a five-second interval.",
    "Long polling was event-rate-dependent. At one event every two seconds, all 100 requests completed roughly 30 times per minute, producing about 3,000 requests per minute.",
    "SSE used 100 initial requests over two minutes, normalized to 49.9 requests per minute, with no reconnect churn during the measured run.",
    "Average delivery delay fell from 2,401.8 ms with short polling to 39.0 ms with long polling and 6.4 ms with SSE.",
    "SSE had the lowest measured average process CPU at 0.87% of one logical core. Long polling had the highest average at 2.57% because it repeatedly completed and recreated requests.",
    "No transport lost an entire client population. The retained raw files include 5,700 short-poll deliveries and 5,900 deliveries for both long polling and SSE."
])

add_section(doc, 16, "Conclusion")
add_body(doc, "For this one-way live notification feed, Server-Sent Events performed best. It delivered notifications with the lowest average delay, used the fewest HTTP requests, and recorded the lowest average Node process CPU. Long polling achieved near-real-time delivery but created the most request churn under the two-second event workload. Short polling was straightforward and produced the expected traffic, but its five-second schedule added multi-second delivery delay.")
add_body(doc, "The benchmark supports the learning objective: choosing the simplest transport that meets the need saves servers and complexity. Short polling remains a valid choice when updates are rare and delay is unimportant. SSE is the simpler efficient choice when a browser needs timely one-way updates. Long polling is a compatibility technique when persistent streaming is not practical.")

add_section(doc, 17, "Future Improvements")
add_bullets(doc, [
    "Add a WebSocket comparison for bidirectional workloads.",
    "Move the notification source to Redis pub-sub so multiple server instances receive the same events.",
    "Test behind Nginx or another reverse proxy and load balancer with realistic timeout and buffering settings.",
    "Run repeated trials and report confidence intervals instead of one controlled run per transport.",
    "Use separate client and server machines so client work does not share the server CPU.",
    "Launch 100 real browser contexts to compare EventSource and browser scheduling with the Node simulator.",
    "Increase client counts gradually to identify saturation points, memory behavior, and reconnect storms.",
    "Add TLS, authentication, per-user filtering, and authorization appropriate for production notifications."
])
add_subheading(doc, "Benchmark Limitations")
add_body(doc, "The tests used local loopback on one Windows machine, and the client simulator shared the host with the server. CPU represents the Node server process, not total system CPU. Each transport was measured once for two minutes, so transient scheduler noise can affect small CPU differences. The Node SSE simulator parsed a streaming fetch response; the React demonstration itself used the required browser EventSource API. Autocannon was intentionally limited to an ordinary HTTP capacity check and was not treated as an SSE or long-poll behavioral test.")

add_subheading(doc, "Reproducibility")
add_body(doc, "Install dependencies with npm install, build with npm run build, run the controlled test with npm run benchmark, run the capacity check with npm run autocannon, and validate deliverables with npm run verify. The results directory preserves the full configuration, raw CPU samples, connection durations, server counters, and per-delivery timestamps used to calculate this report.")

doc.core_properties.title = "FSD Internal 2 Polling vs Long Polling vs Server Sent Events"
doc.core_properties.subject = "Assignment 19 React real-time transport trade-offs"
doc.core_properties.author = ""
doc.core_properties.keywords = "React, Node.js, polling, long polling, SSE, benchmark"
doc.core_properties.comments = "Generated from measured benchmark data and captured project evidence."

doc.save(OUTPUT)
print(f"Created {OUTPUT}")
